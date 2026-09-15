"""Adds Claude's comments to the owner's PID_Controller_Architecture_
Review_20260825.xlsx IN PLACE - appends columns, does not remove or
alter any of the owner's original content. Run after editing the
COMMENTS dicts below.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
PATH = ROOT / "PID_Controller_Architecture_Review_20260825.xlsx"

HEADER_FILL = PatternFill(start_color="2B3A67", end_color="2B3A67", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
AGREE_FILL = PatternFill(start_color="D6EFD6", end_color="D6EFD6", fill_type="solid")
CHANGE_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
NOTAGREE_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
WRAP = Alignment(wrap_text=True, vertical="top")

# (row #, Claude status, Claude comment) - status is my own read on THEIR
# verdict, not a re-vote; "CONFIRMED" = matches what's already built/true,
# "ACTIONABLE NOW" = a concrete next step exists and is cheap, "NOTED FOR
# LATER" = correct but out of scope until the response-surface study.
REVIEW_COMMENTS = {
    1: ("CONFIRMED", "Matches by construction: the adaptive-threshold controller's error term is the bot's own hit-rate, not price. No code anywhere reads market price into a PID error term."),
    2: ("CONFIRMED", "Never built, never proposed after the category-error discussion earlier today. No plant exists for price."),
    3: ("AGREED, DEFERRED", "Correct critique - hit-rate can rise while expectancy stays negative (e.g. cutting big losers AND big winners alike). NOT retrofitted into the frozen META_CONTROL_EXPERIMENT_V1 (freeze discipline: new defect = new candidate version, same as the frozen P02 engine). A V2 outer-loop PV redesign using net expectancy in R is real future work, correctly ordered AFTER the response-surface study per your own Recommended Sequence, not before."),
    4: ("AGREED, DEFERRED", "Same reasoning as #3 - a trend system can be profitable under 40% hit-rate with a fat right tail, so 0.5 as a universal setpoint is arbitrary. Setpoint redesign (expectancy vs. a cost-adjusted hurdle) waits for the same V2 outer-loop work as #3."),
    5: ("CONFIRMED", "Matches exactly what's built: base_threshold + P/I/D adjustment applied to entry_score, the only actuator in the frozen experiment."),
    6: ("CONFIRMED", "No per-study weight adaptation exists anywhere in this codebase. Treating this as a standing constraint - will not build it before incremental-information analysis (Sequence step 4) demonstrates which state dimensions carry independent signal."),
    7: ("AGREED, GAP NOTED", "Real gap in the current (frozen) controller: its P term uses the raw trailing-window hit-rate each step, no additional filtering/smoothing, and gain (kp) is unbounded except via the overall threshold clamp. A future candidate should use a filtered/uncertainty-aware estimate with an explicit gain cap - noted as a V2 design requirement, not retrofitted into the frozen file."),
    8: ("CONFIRMED, ONE GAP", "Anti-windup (hard integral clamp) is implemented and directly tested (test_anti_windup_clamp_bounds_the_integral_contribution). NOT yet implemented: explicit reset/freeze on a structural regime change - the integral currently only decays via ordinary error dynamics, never an explicit reset trigger. Real gap for a V2 design."),
    9: ("CONFIRMED", "Matches the adopted Static->P->PI->filtered-D staged progression. Note: the frozen 36-config grid DID vary kd (0, 2) as part of its coarse sensitivity sweep, but the finding was read as a broad-plateau/no-single-optimum result, not as validating any specific D gain - consistent with 'retain D only if it adds holdout value', which hasn't been separately tested yet."),
    10: ("CONFIRMED", "This is verbatim the backtest script's own stated HONESTY CONSTRAINT: '64 trades is a thin sample... this script draws a directional conclusion... not a final tuned system.'"),
    11: ("CONFIRMED", "Matches the frozen manifest exactly - results.full_grid preserves all 36 rows, and the conclusion field explicitly frames the UNANIMOUS direction of improvement as the credible finding, not any single config."),
    12: ("CONFIRMED", "The -Rs.98 config is never called a winner anywhere in the frozen artifact or its conclusion text."),
    13: ("CONFIRMED", "This is the literal frozen conclusion sentence, word for word."),
    14: ("ACTIONABLE NOW - REAL GAP", "Genuine finding: the current backtest EXCLUDES rejected candidates from any outcome tracking ('rejected trades: no feedback, no P&L counted - never happened' - correct for what the controller itself should learn from, but incomplete for auditing the controller's OWN decisions). Since all 64 V10-C trades already have REAL recorded outcomes (V10-C's own system took all of them), a confusion-matrix breakdown - accepted&won / accepted&lost / rejected-would-have-won / rejected-would-have-lost - is computable RIGHT NOW from existing data, no new acquisition needed. Proposing this as the next concrete build, as a companion diagnostic to (not an edit of) the frozen experiment."),
    15: ("CONFIRMED", "Done exactly as specified: implementation file, trades_csv_sha256, full gain grid, anti-windup rule, window, complete 36-row result set, all in META_CONTROL_EXPERIMENT_V1_FROZEN_20260825.json + .sha256 sidecar."),
    16: ("CONFIRMED", "Matches the standing plan exactly - no PI controller exists; state vector (Layer 2) is being validated first."),
    17: ("CONFIRMED", "Adopted verbatim as the complexity ladder for the eventual inner loop."),
    18: ("NOTED FOR LATER", "Correct and important, but there is no inner-loop controller yet for this to apply to - genuinely nothing to implement today. Recorded as a hard design requirement for whenever Sequence step 7 (Test P then PI) begins."),
    19: ("CONFIRMED", "Matches earlier discussion verbatim: MPC needs a trustworthy state-evolution model we don't have yet; gain scheduling needs learned/validated gains, not invented ones."),
    20: ("CONFIRMED", "This is Layer 2's entire reason for existing - study_layer_v2_state_vector.py, 353 passing tests, S/L_vwap/L_bb/M/dM/V kept as 6 distinct fields + Regime as a 7th separate context dimension, confirmed never fused into one scalar anywhere."),
}

SEQUENCE_COMMENTS = {
    1: ("IN PROGRESS - 14/16 DONE, 2 PARTIAL, ROUND-3 ADDITIONS BUILT",
        "See STUDY_LAYER_V2_REVIEW_TRACKER_20260825.xlsx for the base 16-item gate. "
        "EMPIRICAL FINDING (run against the real, hash-verified V10-C 1-min archive, all 48 symbols): "
        "validate_bar_continuity() found 0/48 symbols fully clean - 388 intra-session gaps, 0 duplicate "
        "timestamps, 0 non-monotonic sequences (chronological integrity: PASS; continuity integrity: PARTIAL). "
        "Gap-size distribution: 1min x128, 2min x55, 3min x96, 6min x10, 23min x1, 28min x2, 90min x96. "
        "The 96 gaps of exactly 90 missing minutes account for 8,640 of 9,305 total missing candles (93%) - "
        "a systematic pattern, not noise. SYMBOL/DATE ATTRIBUTION (completed): the 96 90-min gaps resolve "
        "to exactly 2 CALENDAR DATES - 2024-03-02 and 2024-05-18 - each hitting ALL 48/48 symbols "
        "identically, same exact window 09:59-11:30 (48 x 2 = 96, exact match) - a market-wide or "
        "acquisition-wide event on two specific days, not scattered per-symbol noise. Plus 3 isolated "
        "single-symbol single-day gaps (2023-08-21, 2025-01-06, 2025-12-05) consistent with ordinary "
        "illiquidity, not systemic. RECOMMENDATION: exclude the 2024-03-02 and 2024-05-18 windows (plus "
        "each dimension's own lookback horizon around them, via data_quality_status()) from the forward-"
        "label eligibility set outright, rather than a per-case investigation - the identical cross-symbol "
        "pattern makes this a clean, simple exclusion. Whether the underlying cause was a real exchange-side "
        "halt/DR-drill or an acquisition-side gap is NOT determined here and doesn't need to be for the "
        "purpose of excluding a known-contaminated window from a research eligibility set. "
        "Per the owner's explicit instruction: NO repair applied anywhere (no forward-fill, no interpolation, "
        "no synthetic bars) - detection only, exactly as this module has been built from the start. "
        "BUILT IN RESPONSE (2026-08-25, same day): aggregate_1min_to_5min_with_quality() now reports "
        "source_bar_count / expected_source_bar_count / bar_complete per output bar (a 5-min bar built from "
        "fewer than 5 real 1-min rows is flagged, never silently treated as ordinary); data_quality_status() "
        "classifies each bar as CLEAN / INCOMPLETE_SOURCE_BAR / RECENT_GAP_IN_LOOKBACK, PER DIMENSION, using "
        "that dimension's own real lookback horizon (Structure ~78 bars, Bollinger 20, SMI 10, ATR 14) - "
        "not one global status, since a gap 15 bars back contaminates Ichimoku's cloud but not SMI's momentum "
        "read; vwap_session_contamination_status() handles VWAP separately (SESSION_GAP_UPSTREAM, cumulative "
        "from session start, since VWAP's own accumulation means one early gap contaminates the whole rest of "
        "the session - a genuinely different contamination shape than a fixed lookback window, deliberately "
        "NOT reusing the STALE_CONTEXT label from assert_temporal_alignment, which means something else). "
        "9 new tests, 362/362 passing project-wide. Status per owner's round-3 verdict: "
        "'Layer 2 code architecture: GO / essentially ready. Dataset certification: PARTIAL. "
        "Forward state-information experiment: GO only with gap-aware eligibility.'"),
    2: ("NOT STARTED", "The literal next task once #1's remaining gap is resolved or explicitly bounded."),
    3: ("NOT STARTED", "Depends on #2."),
    4: ("PARTIALLY MOTIVATED", "Today's DISCRETE-vote redundancy measurement (old same-bar Ichimoku vs Bollinger vs VWAP: 78-92% pairwise agreement across 8 liquid symbols, 3 years) already showed real redundancy in the OLD architecture. The analogous measurement on the NEW continuous L_vwap/L_bb state dimensions has not been run yet - a different, still-needed analysis."),
    5: ("NOT STARTED", "Correctly gated behind #2-4 - explicitly not guessed."),
    6: ("NOT STARTED", ""),
    7: ("NOT STARTED", ""),
    8: ("NOT STARTED", ""),
    9: ("NOT STARTED", "Design requirement recorded (see Review row 18); no code yet since there's no controller to constrain."),
    10: ("PARTIALLY DONE", "Frozen experiment exists (hit-rate PV, no counterfactual tracking). Two concrete gaps against this step: (a) counterfactual candidate tracking per Review row 14, (b) expectancy-in-R as PV per Review rows 3-4, both correctly deferred to a V2, never retrofitted into the frozen V1."),
}


def annotate():
    wb = load_workbook(PATH)
    ws1 = wb["PID Controller Review"]
    header_row = 1
    status_col = ws1.max_column + 1
    comment_col = status_col + 1
    ws1.cell(row=header_row, column=status_col, value="Claude - Read on Verdict").fill = HEADER_FILL
    ws1.cell(row=header_row, column=status_col).font = HEADER_FONT
    ws1.cell(row=header_row, column=comment_col, value="Claude's Comments (2026-08-25)").fill = HEADER_FILL
    ws1.cell(row=header_row, column=comment_col).font = HEADER_FONT

    for row_idx in range(2, ws1.max_row + 1):
        n = ws1.cell(row=row_idx, column=1).value
        if n in REVIEW_COMMENTS:
            status, comment = REVIEW_COMMENTS[n]
            sc = ws1.cell(row=row_idx, column=status_col, value=status)
            cc = ws1.cell(row=row_idx, column=comment_col, value=comment)
            sc.alignment = WRAP
            cc.alignment = WRAP
            fill = AGREE_FILL if status == "CONFIRMED" else CHANGE_FILL
            sc.fill = fill

    ws1.column_dimensions[get_column_letter(status_col)].width = 24
    ws1.column_dimensions[get_column_letter(comment_col)].width = 90

    ws2 = wb["Recommended Sequence"]
    status_col2 = ws2.max_column + 1
    ws2.cell(row=1, column=status_col2, value="Claude - Status & Notes (2026-08-25)").fill = HEADER_FILL
    ws2.cell(row=1, column=status_col2).font = HEADER_FONT
    for row_idx in range(2, ws2.max_row + 1):
        n = ws2.cell(row=row_idx, column=1).value
        if n in SEQUENCE_COMMENTS:
            status, note = SEQUENCE_COMMENTS[n]
            text = f"{status}. {note}" if note else status
            cell = ws2.cell(row=row_idx, column=status_col2, value=text)
            cell.alignment = WRAP
    ws2.column_dimensions[get_column_letter(status_col2)].width = 100

    try:
        wb.save(PATH)
        print(f"Annotated and saved {PATH}")
    except PermissionError:
        fallback = PATH.with_name(PATH.stem + "_WITH_CLAUDE_COMMENTS.xlsx")
        wb.save(fallback)
        print(f"{PATH} is locked (likely open in Excel) - saved to {fallback} instead. "
              f"Close the original in Excel and re-run this script to write into it directly.")


if __name__ == "__main__":
    annotate()
