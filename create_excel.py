import pandas as pd
import openpyxl
from openpyxl.chart import LineChart, BarChart, ScatterChart
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import re

# Read the CSV
df = pd.read_csv(r'C:\Users\Dishan\Documents\Codex\Zerodha_live_bot_3.4_ENTRY_UNKNOWN\calibration_progress.csv')

# Rename columns for clarity
df.columns = ['Iteration', 'Win Rate %', 'Trades']

# Calculate rolling averages for trend
df['Win Rate MA (50)'] = df['Win Rate %'].rolling(window=50, center=True).mean()
df['Win Rate MA (100)'] = df['Win Rate %'].rolling(window=100, center=True).mean()

# Calculate statistics
print(f"Total Iterations: {len(df)}")
print(f"Average Win Rate: {df['Win Rate %'].mean():.2f}%")
print(f"Win Rate Std Dev: {df['Win Rate %'].std():.2f}%")
print(f"Min Win Rate: {df['Win Rate %'].min():.2f}%")
print(f"Max Win Rate: {df['Win Rate %'].max():.2f}%")
print(f"Total Trades: {df['Trades'].sum():.0f}")
print(f"Avg Trades per Iteration: {df['Trades'].mean():.1f}")

# Determine phases
phase1_end = 50
phase2_end = 250
df['Phase'] = df['Iteration'].apply(
    lambda x: 'Phase 1: Random' if x <= phase1_end 
    else ('Phase 2: Bayesian' if x <= phase2_end 
    else 'Phase 3: Fine-tuning')
)

# Create Excel workbook
with pd.ExcelWriter(
    r'C:\Users\Dishan\Documents\Codex\Zerodha_live_bot_3.4_ENTRY_UNKNOWN\Calibration_Progress.xlsx',
    engine='openpyxl'
) as writer:
    
    # Write main data
    df.to_excel(writer, sheet_name='Data', index=False)
    
    # Write summary stats
    summary_data = {
        'Metric': [
            'Total Iterations',
            'Average Win Rate %',
            'Win Rate Std Dev %',
            'Min Win Rate %',
            'Max Win Rate %',
            'Total Trades',
            'Avg Trades/Iteration',
            'Phase 1 Iterations (1-50)',
            'Phase 2 Iterations (51-250)',
            'Phase 3 Iterations (251+)'
        ],
        'Value': [
            len(df),
            f"{df['Win Rate %'].mean():.2f}",
            f"{df['Win Rate %'].std():.2f}",
            f"{df['Win Rate %'].min():.2f}",
            f"{df['Win Rate %'].max():.2f}",
            f"{df['Trades'].sum():.0f}",
            f"{df['Trades'].mean():.1f}",
            len(df[df['Iteration'] <= 50]),
            len(df[(df['Iteration'] > 50) & (df['Iteration'] <= 250)]),
            len(df[df['Iteration'] > 250])
        ]
    }
    summary_df = pd.DataFrame(summary_data)
    summary_df.to_excel(writer, sheet_name='Summary', index=False)
    
    # Get workbook to add charts
    workbook = writer.book
    ws_data = writer.sheets['Data']
    
    # Format headers
    for cell in ws_data[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    
    # Add Win Rate Trend Chart
    chart1 = LineChart()
    chart1.title = "Win Rate Progression Over Iterations"
    chart1.style = 12
    chart1.y_axis.title = 'Win Rate %'
    chart1.x_axis.title = 'Iteration'
    chart1.height = 12
    chart1.width = 20
    
    data = openpyxl.chart.Reference(ws_data, min_col=2, min_row=1, max_row=len(df)+1)
    cats = openpyxl.chart.Reference(ws_data, min_col=1, min_row=2, max_row=len(df)+1)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    ws_data.add_chart(chart1, "G2")
    
    # Add Trades Count Chart
    chart2 = BarChart()
    chart2.title = "Trades Generated Per Iteration"
    chart2.style = 11
    chart2.y_axis.title = 'Trades Count'
    chart2.x_axis.title = 'Iteration'
    chart2.height = 12
    chart2.width = 20
    
    data2 = openpyxl.chart.Reference(ws_data, min_col=3, min_row=1, max_row=len(df)+1)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats)
    ws_data.add_chart(chart2, "G20")
    
    # Format Summary sheet
    ws_summary = writer.sheets['Summary']
    for cell in ws_summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

print("\nExcel file created: Calibration_Progress.xlsx")
